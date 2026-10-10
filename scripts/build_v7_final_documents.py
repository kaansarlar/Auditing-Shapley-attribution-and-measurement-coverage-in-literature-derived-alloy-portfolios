"""Merge the v6.2 manuscript with verified completion bounds and exploratory equivalence sensitivity.
Preserves six native OMML equations and four main figures. Uses project-relative paths.
"""
from pathlib import Path
import importlib.util,re,json
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'documents';R=ROOT/'results/v7_final_additions'
spec=importlib.util.spec_from_file_location('format_base',ROOT/'scripts/build_v6_2_documents.py');b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)

def insert(doc,after,text,style='Normal'):
    return b.paragraph_after(doc,after,text,style)

def replace(doc,startswith,text):
    matches=[p for p in doc.paragraphs if p.text.startswith(startswith)]
    assert len(matches)==1,(startswith,len(matches))
    b.replace(matches[0],text);return matches[0]

def save_document(doc,path):
    # Replace the container file completely when a rebuilt ZIP is shorter.
    if path.exists():path.unlink()
    doc.save(path)

def keep_short_tables(doc):
    for table in doc.tables:
        if len(table.rows)<=8:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next=True

def reorder_methods(d):
    body=d._element.body
    children=list(body)
    starts=[(i,e) for i,e in enumerate(children) if e.tag.endswith('}p') and ''.join(e.itertext()).startswith('2.')]
    # Use paragraphs to identify exact top-level Methods headings, avoiding equation text.
    heading_nodes=[p._p for p in d.paragraphs if re.match(r'^2\.\d+ ',p.text)]
    first=children.index(heading_nodes[0]);end=children.index(next(p._p for p in d.paragraphs if p.text=='3 Results'))
    segments=[]
    for j,node in enumerate(heading_nodes):
        start=children.index(node);stop=children.index(heading_nodes[j+1]) if j+1<len(heading_nodes) else end
        segments.append(children[start:stop])
    assert len(segments)==10
    order=[6,7,0,1,2,3,4,5,8,9]
    for e in children[first:end]:body.remove(e)
    for offset,e in enumerate([e for ix in order for e in segments[ix]]):body.insert(first+offset,e)
    for j,p in enumerate([p for p in d.paragraphs if re.match(r'^2\.\d+ ',p.text)],1):
        b.replace(p,re.sub(r'^2\.\d+',f'2.{j}',p.text))
    for p in d.paragraphs:
        if 'defined exactly as in section 2.7' in p.text:b.replace(p,p.text.replace('defined exactly as in section 2.7','defined exactly as in section 2.1'))


def manuscript():
    d=Document(ROOT/'templates/manuscript_v6_2.docx')
    replace(d,'Abstract  ', 'Abstract  Literature-based alloy design seeks promising compositions and an explanation of how individual elements contribute to performance. Yet published measurements cover selected composition series rather than the counterfactual combinations needed to distinguish those contributions. This raises a practical question: can a stable game-theoretic ranking identify elemental contributions, or does it mainly reflect the compositions available in the literature? We address this question for magnetocaloric high-entropy alloys using complete-design synthetic benchmarks and an audited experimental portfolio. Two generators separate measurement and publication effects from errors caused by missing combinations, while grouped resampling tests whether apparent confidence accompanies correct attribution. Within a bounded monotone completion class, a constructive counterexample and linear-program bounds show that the same measured evidence is compatible with every element ranking first or last, even when hypothetical utilities cannot exceed the best measured score. In synthetic portfolios matching the experimental support structure, top-player recovery does not establish an advantage over uniform random selection, including apparently confident subsets. The audited portfolio contains 17 strategies but only four optional-element combinations, with Cr present throughout and Mn always accompanying Al. These constraints prevent separable chemical interpretation of their Shapley allocations. Candidate selection remains useful: two measured alloys retain priority, with an independent fallback, although simple decision rules match Nash-product selection. The resulting workflow distinguishes stable allocation from identified attribution and confines alloy recommendations to the evidence the portfolio can support.')
    p=next(p for p in d.paragraphs if p.text.startswith('Incomplete cooperative games already'))
    b.replace(p,p.text+' We retain known access-envelope values in the constructive example; this is a stronger admissible-completion restriction than retaining measured support scores alone. It should not be confused with experimental knowledge of all subset optima.')
    p=next(p for p in d.paragraphs if p.text.startswith('For each scenario, 500'))
    b.replace(p,p.text.replace('No equivalence margin was prespecified, so no equivalence conclusion is drawn.','No equivalence margin was prespecified. Exploratory 90% interval-inclusion analyses at post hoc margins ±0.05 and ±0.075 follow the two-one-sided-test logic [38] and are reported separately in Supplement S14; they do not replace the primary 95% comparisons or establish confirmatory equivalence.'))
    before=next(p for p in d.paragraphs if p.text=='2.9 Exploratory continuous-grid game')
    node=next(p._p for p in d.paragraphs if p.text.startswith('A corrected independent robustness generator'))
    # Add a distinct Methods subsection before the exploratory grid subsection.
    h=insert(d,node,'2.9 Coalition-preserving completion bounds','Heading 2')
    insert(d,h,'For the audited five-player game, linear programming minimizes and maximizes each exact Shapley value over the 32 coalition variables. We fix v(empty)=0 and the four observed access-envelope values, require monotonicity along every single-player addition, and bound all coalition values by either 1 or the best measured score 0.947268. The Shapley formula is linear in these variables. Separate programs maximize a player’s lead over every other player and its trailing margin below every other player. Every optimized game is reconstructed as a maximum-access envelope with all four measured support scores unchanged, and all 40 endpoint and ranking witnesses are archived. These are global bounds within solver tolerance for the specified bounded monotone class, not grid-search extrema or chemically constrained bounds. Preserving observed coalition maxima is an additional restriction on hypothetical completions; measured support scores alone need not identify those maxima in the unknown complete design. Full constraints, primal–dual checks and independent permutation-based Shapley verification are in Supplement S13.')
    b.replace(before,'2.10 Exploratory continuous-grid game')
    reorder_methods(d)
    p=next(p for p in d.paragraphs if p.text.startswith('The leader changes although'))
    insert(d,p._p,'The broader coalition-preserving completion class yields Shapley ranges of 0–0.813974 for Mn and Al, 0–0.731132 for Cr, 0.002498–0.513463 for Ni and 0.011740–0.582724 for Cu. Every player has an explicit completion making it uniquely first and another making it uniquely last. This remains true when all utilities are capped at 0.947268, the best measured score (Supplementary table S23). Thus no player’s leading or trailing rank is identified even in this more restricted class. The endpoints are elementwise bounds; they are not generally attained simultaneously by one game.')
    p=next(p for p in d.paragraphs if p.text.startswith('Exact calculation, resampling stability'))
    b.replace(p,p.text+' The post hoc interval-equivalence sensitivity supports margins of ±0.075 for the full exact-geometry sets, but not for the small confident independent-generator subsets; conclusions at ±0.05 are less consistent (Supplementary table S24). These exploratory margins have no established materials-specific meaning.')
    p=next(p for p in d.paragraphs if p.text.startswith('Both generators use five binary'))
    b.replace(p,p.text.replace('No practical-equivalence margin or confirmatory test of benefit from bootstrap confidence was prespecified.','Practical-equivalence margins were chosen post hoc; no confirmatory equivalence analysis or paired test of benefit from bootstrap confidence was prespecified. Completion bounds assume bounded monotonicity and preservation of four observed access-envelope values. They omit chemical feasibility, processing constraints and additive or convex structure; a particular scientific assumption could narrow them.'))
    p=next(p for p in d.paragraphs if p.text.startswith('Incomplete-support Shapley ambiguity is'))
    b.replace(p,p.text.replace('A constructive pair of complete designs retains the same measured evidence but changes the Shapley leader.','A constructive pair retains the same measured evidence but changes the Shapley leader; coalition-preserving completion bounds show that every player can be uniquely first or last, also under the best-measured-score cap.').replace('they do not demonstrate equivalence.','the primary comparisons do not demonstrate equivalence.'))
    # Preserve existing external hyperlinks in Data availability by editing text runs only.
    p=next(p for p in d.paragraphs if p.text.startswith('The source tables are available'))
    for run in p.runs:run.text=run.text.replace('v6.2','v7 final').replace('verified completion witness and decision controls','verified completion witnesses, global coalition-preserving bounds, exploratory equivalence sensitivity and equal-budget decision controls')
    p=next(p for p in d.paragraphs if p.text.startswith('The accompanying v6.2 archive'))
    b.replace(p,'The accompanying v7 final archive uses project-relative paths and fixed seeds. It includes scenario-level results, separate generator definitions, exact random baselines, uncertainty intervals, conditional confidence denominators, six editable Word equations, integrity hashes and tests. Completion bounds and exploratory sensitivity are regenerated with scripts/run_v7_final_completion_bounds.py and scripts/run_v7_final_equivalence.py. The witness and equal-budget controls use scripts/run_v6_2_additions.py. Existing long benchmark outputs are frozen rather than claimed as newly simulated; new analyses operate on those archived records.')
    d.add_paragraph('[38] Lakens D 2017 Equivalence tests: a practical primer for t tests, correlations, and meta-analyses Soc. Psychol. Personal. Sci. 8 355–362 https://doi.org/10.1177/1948550617697177')
    # Every main figure and table is explicitly discussed in running text.
    links=[('We quantify that problem','Figure 1 summarizes the two linked analysis paths.'),
           ('Write q(P)','Table 1 compares the resulting exact allocations.'),
           ('Across 4000 worlds','Figure 2 shows the generator-specific comparisons and exact-geometry recovery.'),
           ('Top ties and wrong unique selections','Table 2 reports the paired material-error comparisons.'),
           ('The 25 strict 3d-transition-metal','Figure 3 locates the measured candidates in the temperature–entropy plane.'),
           ('Cr occurs in all 17','Table 4 lists the four support classes and their measured maxima.'),
           ('G8 contains S30','Figure 4 combines support constraints with research-group omission results.')]
    for prefix,sentence in links:
        para=next(p for p in d.paragraphs if p.text.startswith(prefix));b.replace(para,para.text+' '+sentence)
    heading=next(p for p in d.paragraphs if p.text.startswith('3.3 The audit'))
    audit=insert(d,heading._p,'The target audit changes which observations can support each analysis, rather than simply reducing the sample size. Table 3 distinguishes the 199 linked records used to trace data lineage from the 141 unique composition-temperature records used as the audit basis. Neither total is the effective sample size for every endpoint: the strict Curie analysis retains 72 records, whereas the primary field-restricted entropy analysis retains 123. The strict 3d-transition-metal decision family contains only 25 records, so predictive evidence from the broader literature cannot automatically be transferred to the candidate-selection problem.')
    insert(d,audit,'These differences follow from the measurement distinctions documented in section 2.3. An entropy-peak temperature cannot be treated as a Curie temperature without supporting evidence, and entropy estimates inferred from a single magnetization-temperature curve are excluded from the comparable isothermal endpoint. Correcting the reported field also changes eligibility under the primary field threshold; lowering that threshold to 0.5 T expands the entropy cohort to 136 records as a sensitivity analysis. Likewise, the 71 phase-annotated records support a separate phase screen rather than a complete annotation of the decision portfolio.')
    caption=next(p for p in d.paragraphs if p.text.startswith('Table 3.'))
    insert(d,caption._p,'The resulting evidence is endpoint-specific and partly overlapping. The cohorts should therefore not be added together or treated as interchangeable validation samples. Temperature prediction, field-normalized entropy screening and phase screening are evaluated on their own eligible records, while measured candidate comparisons remain confined to the strict decision family. This separation is the practical consequence of the audit: a larger linked dataset supplies provenance and broader validation opportunities, but does not by itself increase the counterfactual support available for elemental attribution.')
    b.format_doc(d)
    keep_short_tables(d)
    for para in d.paragraphs:
        if re.match(r'^\[\d+\]',para.text):
            para.paragraph_format.space_after=Pt(3)
            for run in para.runs:run.font.size=Pt(9.5)
    save_document(d,OUT/'MLST_HEA_manuscript_v7_final.docx')


def supplement():
    d=Document(ROOT/'templates/supplement_v6_2.docx')
    for p in d.paragraphs:
        if 'v6.2 package' in p.text:b.replace(p,p.text.replace('v6.2 package','v7 final package'))
        if p.text.startswith('All paired comparison intervals'):
            b.replace(p,p.text.replace('No practical-equivalence margin was prespecified.','No practical-equivalence margin was prespecified; Supplement S14 reports an explicitly post hoc sensitivity analysis.'))
    d.add_heading('S13 Coalition-preserving completion bounds',level=1)
    d.add_paragraph('The main counterexample fixes measured support scores and four observed access-envelope values. We retain this stronger restriction for a bounded monotone completion class, while recognizing that the complete physical landscape need not preserve those observed maxima: an unknown subset could improve them. Let O={Cr–Ni, Cr–Ni–Cu, Mn–Cr–Al, Mn–Cr–Ni–Al}, with values from main Table 1. The feasible class fixes v(empty)=0, v(S) to the measured access value for S in O, and imposes 0≤v(C)≤M and v(C)≤v(C∪{i}) for every absent player i. We analyze M=1 and M=0.9472682371859096. The latter also fixes the grand-coalition maximum to the best measured score and is an additional scenario assumption.')
    d.add_paragraph('The exact Shapley formula (main equation 6) supplies a linear coefficient vector for each player. Minimizing and maximizing that linear objective gives its global range over the feasible class. To construct a unique leader i, maximize t subject to φᵢ−φⱼ≥t for all j≠i; to construct a unique last player, maximize t subject to φⱼ−φᵢ≥t. A positive optimum supplies a strict ranking witness. All programs use scipy.optimize.linprog with the HiGHS solver. Primal feasibility, primal–dual agreement within 10⁻⁸, boundedness, monotonicity, observed-value agreement and Shapley efficiency are checked. Shapley values are also independently computed by averaging marginal contributions over all 120 player permutations.')
    d.add_paragraph('Every feasible monotone v is realized in this finite support class by setting q(C)=v(C) at unobserved supports, q(S) to its measured score at observed S, and q(empty)=0. Here v(S)=q(S) is fixed at all four observed supports, so max{q(P):P⊆C}=v(C). This construction is checked numerically for every witness. It holds for the stated scalar bounded-access model and does not assert realizability of a physical alloy or positivity of all Möbius coefficients.')
    data=pd.read_csv(R/'coalition_preserving_shapley_bounds.csv')
    rows=[]
    for player in ['Mn','Cr','Ni','Al','Cu']:
        a=data[(data.player==player)&(data.cap_class=='unit_cap')].iloc[0]
        c=data[(data.player==player)&(data.cap_class=='measured_cap')].iloc[0]
        rows.append([player,f'{max(0,a["min"]):.6f}–{a["max"]:.6f}',f'{max(0,c["min"]):.6f}–{c["max"]:.6f}',f'{c.lead:.6f}',f'{c.trail:.6f}','Yes / yes'])
    node=d.paragraphs[-1]._p
    table=b.table_after(d,node,['Player','Range M=1','Range M=0.947268','Lead margin at cap','Trail margin at cap','Unique top / last at cap'],rows,[.5,1.25,1.3,1.0,1.0,1.0])
    insert(d,table._tbl,'Table S23. Global Shapley bounds within the coalition-preserving bounded monotone class, and attained ranking margins under the best-measured-score cap. Each endpoint and margin has its own completion; elementwise extremes need not be jointly attainable. Mn and Al have equal range bounds because their observed constraints are symmetric, although a particular completion may break their equivalence.')
    d.add_paragraph('The archive supplies coalition_preserving_completion_witnesses.csv with all 32 support utilities, all 32 access values, five exact Shapley values and primal–dual objectives for every one of the 40 solutions. The script run_v7_final_completion_bounds.py regenerates the witnesses and bounds. These bounds replace any reliance on an unrecorded multistart-grid witness; they refer to the stricter coalition-preserving class and should not be substituted for bounds under a support-score-only completion class.')
    d.add_heading('S14 Exploratory equivalence sensitivity',level=1)
    d.add_paragraph('The primary claims use 95% difference and recovery intervals. A failure to exclude zero or recovery 0.20 is not proof of equivalence. Following the interval-inclusion logic of two one-sided tests (main reference [38]), this exploratory sensitivity asks whether a 90% interval lies strictly inside a chosen equivalence band. We use ±0.05 and ±0.075 for material-error differences, and 0.20±0.05 and 0.20±0.075 for exact recovery. These margins were chosen after inspecting results and have no established materials-specific interpretation. They are not prespecified confirmatory tests.')
    d.add_paragraph('Material-error differences use 4000 paired scenario bootstrap draws with seed 20261008. Exact recovery uses 90% Wilson intervals. Exact-recovery denominators exclude tied complete-design truth; material-error calculations retain these worlds when the estimated top is unique. Confidence-subset counts are recomputed from the archived records rather than copied as constants. Primary confidence concerns the unique bootstrap modal player; independent confidence concerns the estimated top frequency at the independent threshold. These generator-specific definitions are retained without pooling.')
    eq=pd.read_csv(R/'exploratory_equivalence_sensitivity.csv');rows=[]
    def yes(v):return 'Yes' if v==True or v=='True' else 'No'
    for row in eq.itertuples():
        m='—' if row.material_n==0 else f'{row.material_difference:+.4f} [{row.material_90_low:+.4f}, {row.material_90_high:+.4f}]'
        er=f'{row.exact_k}/{row.exact_n} [{row.exact_90_low:.4f}, {row.exact_90_high:.4f}]'
        series=eq[(eq.generator==row.generator)&(eq.subset==row.subset)].iloc[0]
        flags=[]
        for delta in [.05,.075]:
            flags.append(('—' if row.material_n==0 else yes(series[f'material_within_{delta}']))+' / '+yes(series[f'exact_within_{delta}']))
        rows.append([row.generator,row.subset,m,er,*flags])
    table=b.table_after(d,d.paragraphs[-1]._p,['Generator','Subset','Error difference [90% CI]','Recovery k/n [90% CI]','Within ±0.05 error / recovery','Within ±0.075 error / recovery'],rows,[.7,1.4,1.3,1.25,.8,.8])
    insert(d,table._tbl,'Table S24. Post hoc interval-equivalence sensitivity. Error denominators are 821, 339 and 358 in the primary generator and 804, 583 and 288 in the independent generator; recovery denominators are shown explicitly. Confidence-only rows assess recovery. A band can contain a 90% interval even if that interval excludes zero; a small detected difference and equivalence within a wide margin can coexist.')
    d.add_paragraph('At ±0.075, both full HEA-matched sets meet interval-inclusion conditions for error differences and recovery, and both veto-only sets do so. Veto-plus-equivalence error differences fail in both generators, and the independent combined class also fails for recovery. The small confident independent subsets do not establish recovery equivalence at either margin. At ±0.05 conclusions vary across rows. These results quantify sensitivity to a chosen tolerance; they do not establish that bootstrap confidence has no benefit or a universal random-equivalent attribution rule.')
    # Discuss every supplementary display in running text, beyond its caption.
    links=[('The v3 master table','Table S1 summarizes the target and composition policies.'),
           ('The prespecified 250-350 K','Table S10 reports candidate changes across the translated windows, and Figure S2 separates those changes from structural allocation invariances.'),
           ('The continuous grid uses','Figure S3 displays the exploratory grid game, and Table S11 reports its sensitivity to publication omissions.'),
           ('Each synthetic world contains','Table S12 specifies the primary simulation design.'),
           ('A bootstrap result is confident','Table S13 summarizes these outcomes.'),
           ('The 20 full-coverage','Table S15 relates estimation and structural agreement to main-effect rank.'),
           ('The primary 123-record','Table S21 summarizes the predictive gates and Figure S4 shows the underlying endpoint diagnostics.'),
           ('For 20000 Dirichlet','Table S22 compares regret at the same selection budgets.'),
           ('The main counterexample fixes','Table S23 reports the resulting bounds and strict-ranking margins.'),
           ('Material-error differences use','Table S24 gives every interval, recovery denominator and sensitivity classification.')]
    for prefix,sentence in links:
        para=next(p for p in d.paragraphs if p.text.startswith(prefix));b.replace(para,para.text+' '+sentence)
    headings=[('S2 Predictive validation','Regression transfer and direct temperature-window classification are evaluated separately in Tables S2 and S3; a global signal does not establish reliable transfer within the intended alloy family.'),
        ('S3 Empirical strategy portfolio','Table S4 lists the measured candidates in score order, while Table S5 groups them by optional-element support. Table S6 gives the quotient-game allocation when the support-equivalent Mn and Al players are merged.'),
        ('S4 Research group and field sensitivity','Table S7 shows which measured candidate remains preferred after research-group removal or the field restriction; utilities are recomputed within each retained portfolio.'),
        ('S5 Null analyses and robustness','Table S8 tests target alignment while conditioning on element incidence. Table S9 records candidate frequencies across utility settings, and Figure S1 compares preference regret for fixed decision rules.'),
        ('S9 Additional benchmark analyses','Table S17 compares unique-top material error with its paired random baseline, and Table S18 reports exact-geometry recovery. Tables S19 and S20 distinguish held-out diagnostic performance from corrected independent-generator attribution results.')]
    for prefix,sentence in headings:
        heading=next(p for p in d.paragraphs if p.text.startswith(prefix));insert(d,heading._p,sentence)
    for table_id,sentence in [('Table S14.','Table S14 reports the held-out AUC loss after permuting each diagnostic; the largest loss occurs for the estimated top margin.'),
         ('Table S16.','Table S16 summarizes all 600 HEA-like stress worlds. This broader set includes designs that do not exactly reproduce the observed support geometry and therefore has a different denominator from Table S18.')]:
        cap=next(p for p in d.paragraphs if p.text.startswith(table_id));table_node=cap._p.getprevious()
        para=d.add_paragraph(sentence);table_node.addprevious(para._p)
    # Remove inherited hard page breaks that stranded a heading above its table.
    for para in list(d.paragraphs):
        for br in para._p.xpath('.//w:br[@w:type="page"]'):
            br.getparent().remove(br)
        if not para.text and not para._p.xpath('.//w:drawing'):
            # Preserve image paragraphs; remove only now-empty break paragraphs.
            if para._p.getnext() is not None and para._p.getnext().tag.endswith('}tbl'):
                para._p.getparent().remove(para._p)
    table=next(t for t in d.tables if t.cell(0,0).text=='Rank')
    table.cell(0,0).text='Order'
    replace(d,'Table S4.','Table S4. All 17 eligible observed strategies in descending Nash-score order. Exact score ties retain source-row order; S30 and S32 are co-optimal rather than distinct first and second ranks.')
    cap=next(p for p in d.paragraphs if p.text.startswith('Figure S2.'))
    prev=cap._p.getprevious()
    from docx.text.paragraph import Paragraph
    pic=Paragraph(prev,d);pic.clear();pic.add_run().add_picture(str(ROOT/'figures/Figure_S2_window_sensitivity_v7_final.png'),width=Inches(6.5))
    b.format_doc(d);keep_short_tables(d);save_document(d,OUT/'MLST_HEA_supplement_v7_final.docx')


def revision_summary():
    d=Document();d.add_heading('Revision summary for the final HEA manuscript',0)
    d.add_paragraph('This revision combines the v6.2 audited manuscript with verified completion and equivalence analyses. It preserves the measured candidate hierarchy and narrows elemental interpretation to what the observed portfolio can support. The summary records the scientific changes and their validation.')
    points=[('Completion definition and ties','The main two-design counterexample preserves every measured score and all four observed access-envelope values. Exact Mn–Al symmetry is preserved in the observed game. No artificial ordering of equal Shapley values is used.'),('Global completion bounds and witnesses','The multistart-grid concept is strengthened to global linear-program bounds within a stated coalition-preserving bounded monotone class. All 40 optimized games and their 1280 support rows are recorded, with independent permutation Shapley verification and primal–dual checks. Every player can be uniquely first or last, also at the best-measured-score cap.'),('Statistical language','Primary 95% comparisons remain failure-to-establish-difference statements. A separate post hoc 90% interval-equivalence sensitivity uses ±0.05 and ±0.075. Confident independent subsets do not establish equivalence.'),('Nash selection controls','The v6.2 equal-budget controls, declared source-row tie rule and exact random-subset expectations are retained. S30 and S32 remain co-maximizers of the clipped equal-weight score; a one-candidate deterministic control selects S30 under the common tie rule. Nash-specific superiority is not claimed.'),('Structure and scope','Synthetic methods are presented first. Four main figures and six native editable Word equations are retained. Detailed model and continuous-grid diagnostics remain supplementary. The title and conclusions remain specific to magnetocaloric HEA portfolios and the two simulated generators.'),('Validation and availability','New analyses and equal-budget controls are regenerated; frozen long synthetic runs are retained without claiming a fresh rerun. The package contains source data, scripts, tests, full-precision witnesses, documents and a SHA-256 manifest. No external repository is modified and no new release DOI is claimed.')]
    for title,txt in points:d.add_heading(title,1);d.add_paragraph(txt)
    b.format_doc(d)
    for s in d.sections:s.top_margin=Inches(.8);s.bottom_margin=Inches(.8)
    d.styles['Normal'].font.size=Pt(10.5)
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(20)
    d.styles['Title'].paragraph_format.space_after=Pt(12)
    d.styles['Heading 1'].font.size=Pt(12)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    for style in d.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for para in d.paragraphs:
        for border in para._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    save_document(d,OUT/'Revision_summary_v7_final.docx')

if __name__=='__main__':
    OUT.mkdir(exist_ok=True);manuscript();supplement();revision_summary()
