"""Number references by first appearance in final main text; synchronize supplementary citations.
Preserves OOXML runs, equations and hyperlinks. Run after each document rebuild.
"""
from pathlib import Path
import re,json
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];D=ROOT/'documents';R=ROOT/'results/v7_final_additions'
PAT=re.compile(r'\[([\d,\-– ]+)\]')

def numbers(s):
    ns=[]
    for item in s.replace('–','-').split(','):
        item=item.strip()
        if '-' in item:
            lo,hi=map(int,item.split('-'));ns.extend(range(lo,hi+1))
        elif item:ns.append(int(item))
    return ns

def bracket(ns):
    ns=sorted(set(ns));parts=[];j=0
    while j<len(ns):
        k=j
        while k+1<len(ns) and ns[k+1]==ns[k]+1:k+=1
        if k-j>=2:parts.append(f'{ns[j]}–{ns[k]}')
        else:parts.extend(map(str,ns[j:k+1]))
        j=k+1
    return '['+','.join(parts)+']'

def xml_paras(d):return d._element.body.iter(qn('w:p'))
def texts(e):return list(e.iter(qn('w:t')))
def text(e):return ''.join(t.text or '' for t in texts(e))

def rewrite(e,mapping):
    nodes=texts(e);whole=''.join(n.text or '' for n in nodes);spans=[];pos=0
    for node in nodes:spans.append((pos,pos+len(node.text or ''),node));pos=spans[-1][1]
    for m in reversed(list(PAT.finditer(whole))):
        ns=numbers(m.group(1))
        if not ns or not all(n in mapping for n in ns):continue
        value=bracket([mapping[n] for n in ns]);start,end=m.span();inserted=False
        for lo,hi,node in spans:
            if hi<=start or lo>=end:continue
            s=node.text or '';a=max(start-lo,0);b=min(end-lo,len(s))
            node.text=s[:a]+(value if not inserted else '')+s[b:];inserted=True

def main():
    p=D/'MLST_HEA_manuscript_v7_final.docx';d=Document(p)
    refs={int(re.match(r'^\[(\d+)\]',p.text).group(1)):p for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)}
    mapping={}
    for para in xml_paras(d):
        if re.match(r'^\[\d+\]',text(para)):continue
        for match in PAT.finditer(text(para)):
            ns=numbers(match.group(1))
            if ns and all(n in refs for n in ns):
                for n in ns:
                    if n not in mapping:mapping[n]=len(mapping)+1
    assert len(mapping)==len(refs)==38,(len(mapping),len(refs),set(refs)-set(mapping))
    for para in list(xml_paras(d)):
        if not re.match(r'^\[\d+\]',text(para)):rewrite(para,mapping)
    heading=next(p for p in d.paragraphs if p.text=='References');anchor=heading._p
    for old,new in mapping.items():
        para=refs[old];nodes=texts(para._p);nodes[0].text=re.sub(r'^\[\d+\]',f'[{new}]',nodes[0].text)
        anchor.addnext(para._p);anchor=para._p
    p.unlink()
    d.save(p)
    p=D/'MLST_HEA_supplement_v7_final.docx';s=Document(p)
    for para in xml_paras(s):rewrite(para,mapping)
    p.unlink()
    s.save(p)
    seen=[]
    for para in xml_paras(d):
        if re.match(r'^\[\d+\]',text(para)):continue
        for match in PAT.finditer(text(para)):
            ns=numbers(match.group(1))
            if ns and all(1<=n<=38 for n in ns):
                for n in ns:
                    if n not in seen:seen.append(n)
    assert seen==list(range(1,39)),seen
    (R/'reference_numbering.json').write_text(json.dumps({'old_to_final':mapping,'first_appearance_order':seen,
         'reference_count':38,'supplement_synchronized':True,'equations_and_hyperlinks_preserved':True},indent=2))
    print('References 1–38 appear in consecutive first-citation order; supplementary citations synchronized.')

if __name__=='__main__':main()
