"""Build v6.2 from frozen v6.1 templates while preserving native Word equations."""
from pathlib import Path
from copy import deepcopy
import json,re
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'documents'
TITLE='Diagnosing Shapley identifiability in magnetocaloric high-entropy alloy portfolios'
FAIR=json.loads((ROOT/'results/v6_fairness/v6_fairness_summary.json').read_text())

def replace(p,text):
    p.clear();r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(10.5 if p.style.name=='Normal' else 12)
    r.font.color.rgb=RGBColor(0,0,0)

def paragraph_after(doc,node,text,style='Normal'):
    p=doc.add_paragraph(text,style=style);node.addnext(p._p);return p._p

def table_after(doc,node,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';node.addnext(t._tbl)
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=str(x)
    if widths:
        t.autofit=False
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Inches(w)
    return t

def fair_rows():
    return [[r['cls'].capitalize(),r['n'],f"{r['material_error']:.3f}",f"{r['random']:.3f}",
             f"{r['difference']:+.3f} [{r['diff_ci_low']:+.3f}, {r['diff_ci_high']:+.3f}]"]
            for r in FAIR['primary_fair_baseline_unique_top'][:4]]

def fill(t,rows):
    assert len(t.rows)==len(rows)+1
    for rr,values in zip(t.rows[1:],rows):
        for c,x in zip(rr.cells,values):c.text=str(x)

def format_doc(d):
    for style in d.styles:
        if style.type==1:
            style.font.color.rgb=RGBColor(0,0,0)
    for p in d.paragraphs:
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.bold=True
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    for t in d.tables:
        if t._tbl.xpath('.//m:oMath'):continue
        t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for ri,row in enumerate(t.rows):
            trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
            if ri==0:trp.append(OxmlElement('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=c._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                tcpr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBF0' if ri==0 else ('F8F9FA' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
                tcpr.append(margins)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)

def hyperlink(p,label,url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);r.append(pr)
    text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);p._p.append(h)

def main_document():
    d=Document(ROOT/'templates/manuscript_v6_1.docx');p=list(d.paragraphs);t=list(d.tables)
    replace(p[0],TITLE);p[0].style='Title'
    replace(p[4], 'Abstract  Literature portfolios can support candidate selection without identifying separate element contributions. We study this distinction for magnetocaloric high-entropy alloys using two synthetic generators and an audited Fe–Co-background case. Complete five-player designs provide exact Shapley truth; noise-free sampled designs isolate missing-support error from measurement and publication effects. A constructive example preserves every observed score and coalition value but reverses the leading player after one unobserved combination is completed. Across 4000 primary worlds, median rank agreement is 0.900 between noisy and noise-free games on identical supports but 0.580 between sampled and complete designs. With a veto player alone, material error is 0.302 versus 0.327 under uniform random selection; the corrected independent generator gives 0.632 versus 0.657. Neither paired interval establishes a difference, and neither demonstrates equivalence. With veto and player equivalence together, the direction changes between generators. Among 358 worlds matching the audited support geometry, exact top recovery is 0.209 [0.171, 0.255] and remains 0.207 and 0.202 in apparently confident grouped and publication-bootstrap subsets. These intervals do not establish recovery above the five-player baseline 0.20. The real portfolio contains 17 strategies in four support classes, with Cr as a veto player and Mn–Al support equivalence. S30 and S32 remain measured priorities, with an independent S36 fallback, but their element allocations are not chemical rankings. Equal-budget controls show no Nash-specific advantage over simple deterministic rules. The contribution is a reproducible diagnostic benchmark that distinguishes allocation stability, incomplete-support identification and evidence-constrained selection.')
    replace(p[9], 'Incomplete cooperative games already formalize unknown coalition values and analyze admissible extensions and allocation ranges [36,37]. Their distinction between unavailable information and forbidden cooperation is essential here. A literature portfolio is a collection of measured alloy strategies, not a complete characteristic function for all physically possible compositions. Our accessible-strategy game is fully specified by that collection, whereas the complete performance landscape remains unknown. We therefore do not claim to discover Shapley non-identifiability or to solve the extension problem. Existing cautions about value-function choice and off-manifold explanations also apply [33,34]. What requires empirical assessment is how selective literature coverage and clustered records affect apparent allocation confidence and candidate decisions [32].')
    replace(p[10], 'We quantify that problem using complete-design truth, a noise-free sampled reference, two independently designed generators, matched uniform-random losses and publication- versus research-group bootstrap. Observable support diagnostics are evaluated on held-out worlds rather than assumed to certify an individual rank. The audited magnetocaloric HEA portfolio then illustrates a diagnosed support class. The new evidence concerns the interaction of missing combinations, study selection and conditional false confidence, and the separation of attribution failure from regret among measured candidates. The benchmark supplies a reusable analysis protocol, not a new Shapley allocation formula or an experimentally established alloy discovery.')
    replace(p[18],p[18].text+' Standard composition-based descriptor conventions motivate this representation [13-15].')
    replace(p[30], 'With a zero disagreement vector and equal objective weights, we maximize the Nash product on the finite measured candidate set and report its cube root [19]. This is a discrete Nash-product selection rule, not a noncooperative Nash equilibrium. We do not convexify the attainable objective set or allow randomized mixtures of physically distinct alloys. Consequently, the classical bargaining theorem and a unique bargaining solution are not claimed; all co-maximizers are retained. The choice depends on the utility transforms, clipping thresholds and disagreement point.')
    replace(p[31], 'Let P(a) denote the optional-element support of strategy a and 𝒫 = {Mn, Cr, Ni, Al, Cu}. The characteristic value for coalition C is the highest empirical Nash score among measured strategies with P(a) contained in C; an empty accessible set has value zero. This zero denotes absence of access to observed evidence, not zero magnetocaloric performance of every unmeasured C-supported alloy. The resulting portfolio game is fully defined on all 32 coalitions. It is distinct from an incomplete game for the unknown complete performance landscape [36,37]. Exact Shapley values and interaction indices allocate marginal access [20,21], not causal element effects or feature-SHAP values [22]. Veto players, equivalent incidence columns, support classes and direct contrasts are diagnosed before interpretation. A Mn+Al quotient game and supplementary Banzhaf values provide descriptive checks [28].')
    replace(p[34], p[34].text+' For equal-budget decision controls, 20000 Dirichlet(1,1,1) preference vectors weight the three transformed objectives in a geometric utility. The oracle is the best measured candidate for each preference. Regret is oracle utility minus the best utility in the selected list, at budgets one, three and five. Nash, equal arithmetic, maximin, single-objective, window-then-raw-entropy and uniform-random-subset controls share the same candidate set. Deterministic score ties use original source-row order, not an inherited Nash ordering. Random-subset regret is evaluated by exact order-statistic probabilities. These preferences are a diagnostic assumption, not elicited stakeholder choices; the geometric oracle favors the same objective geometry as the Nash score.')
    replace(p[38],p[38].text+' The intercept-plus-main-effect rank diagnoses linear coefficient separability only. Full rank does not identify unrestricted interaction terms or guarantee recovery of complete-design Shapley values.')
    replace(p[39],p[39].text+' Paired differences and Wilson intervals are uncertainty summaries. An interval containing zero or 0.20 is reported as failure to establish a difference or above-baseline recovery, not evidence of practical equivalence. No equivalence margin was prespecified, so no equivalence conclusion is drawn.')
    # Keep model and grid diagnostics available, but move detailed plots and model table to the supplement.
    replace(p[43], 'Continuous-composition modeling is not part of the identification claim. Its methods, conditional Gaussian-process uncertainty [27] and publication-omission outcomes are reported solely as an exploratory analysis in Supplement S7.')
    p[40].style='Heading 2';p[42].style='Heading 2'
    replace(p[63],p[63].text+' Even the full-rank median of 0.700 is below perfect identification. Matrix rank is therefore a risk descriptor for a specified main-effect model, not a sufficient criterion for interacting Shapley recovery.')
    replace(p[64], 'Top ties and wrong unique selections are reported separately. A unique estimated top is absent in 1.32% of primary worlds. Without veto or equivalence, material error is 0.094 versus 0.326 under random selection; with equivalence alone it is 0.149 versus 0.323. With veto alone it is 0.302 versus 0.327, difference −0.025 [−0.053, +0.004]. The corrected independent generator gives 0.632 versus 0.657, difference −0.026 [−0.056, +0.005]. Neither interval establishes a difference; this does not prove equivalence. With veto plus equivalence, the primary difference is +0.045 [+0.002, +0.089], but the corrected independent difference is −0.064 [−0.101, −0.028]. The direction for the combined class is generator-sensitive.')
    replace(p[66], 'Of 600 HEA-like primary worlds, 358 match the observed support geometry exactly. Every estimated top is the veto player. Material error is 0.304 versus 0.327 under random selection, difference −0.022 [−0.064, +0.021]. Exact recovery is 75/358=0.209 [0.171, 0.255]. Among apparently confident results, recovery is 52/251=0.207 [0.162, 0.262] for group bootstrap and 57/282=0.202 [0.159, 0.253] for publication bootstrap. These intervals do not establish recovery above the uniform five-player benchmark 0.20, including within confident subsets; they are not a paired test that confidence has no benefit. The corrected independent targeted run has 288 exact matches among 9000 worlds. Its material error is 0.653 versus 0.639, difference +0.014 [−0.037, +0.063]. Exact recovery is 44/270=0.163 [0.124, 0.212] after excluding 18 tied-truth worlds, and 14/66=0.212 [0.131, 0.325] and 13/69=0.188 [0.114, 0.296] in publication- and group-confident subsets. None establishes above-baseline recovery. Absolute error rates remain generator-specific.')
    fill(t[9],fair_rows())
    replace(p[69], 'Figure 2. Generator-specific random-baseline analysis. (a) Primary material error among unique-top worlds. (b) Exact top recovery with 95% Wilson intervals in the 358-world primary exact-geometry subset; the dashed line is uniform recovery 0.20. (c) Corrected independent-generator material error. Veto-only difference intervals include zero in both generators, without demonstrating equivalence. The combined veto-and-equivalence difference changes sign.')
    p[68].clear();p[68].add_run().add_picture(str(ROOT/'results/v6_1_postprocessing/Figure_4_v6_1_generator_comparison.png'),width=Inches(6.7))
    replace(p[53],p[53].text+' Both co-optima have temperature and clipped entropy utilities of one and reliability utility 0.85; thus both scores equal the cube root of 0.85, or 0.947268. The tie is a consequence of those utility transformations, not equality of their raw entropy changes or independent corroboration.')
    replace(p[54], 'At budget one, Nash mean regret is 0.000223 and 99.84% of preference draws lie within 0.01 of the measured oracle. Equal arithmetic, maximin, temperature-only, entropy-only and window-then-raw-entropy rules give exactly the same result because each selects S30 under the declared tie rule. The uniform-random one-candidate control has expected regret 0.316128. At budget five the Nash list has zero regret; arithmetic and the raw-entropy window control also achieve zero regret. These internal comparisons support selection among measured candidates but do not establish a Nash-specific advantage. They depend on the same utility transforms and a geometric preference oracle. Formula-derived and source-block representations leave the empirical game unchanged because it uses measured payoffs and player presence. Full equal-budget results and deterministic selections are in Supplement S11.')
    replace(p[59],p[59].text+' The rank deficiency concerns linear main effects, whereas the demonstrated veto and symmetry constraints follow directly from the game definition; neither rank nor their absence is a general Shapley-identification theorem.')
    replace(p[56],p[56].text.replace('Figure 3.','Figure 3.'))
    replace(p[74],p[74].text.replace('Figure 5.','Figure 4.'))
    replace(p[88],p[88].text+' The constructive completion witness shows this is not merely a numerical noise problem. The complete-design attribution changes while the currently observed evidence does not. It is an existence argument for non-identification under bounded utility, not evidence that a high-performing Mn-only alloy physically exists.')
    replace(p[90], 'Exact calculation, resampling stability and recovery of complete-design attribution are different properties. With veto alone, neither generator establishes material-error performance different from the paired uniform-random baseline; this is not an equivalence finding. With veto and equivalent players together, the direction changes between generators. In matched-support subsets, overall and apparently confident top-recovery intervals include 0.20 and do not establish recovery above that benchmark. Bootstrap resamples the evidence already present and cannot add missing counterfactual combinations. The conditional wrong rates quantify this limitation but do not establish that bootstrap confidence causes or increases error. Candidate regret is a separate quantity and must be judged against equal-budget controls.')
    replace(p[94], 'A confirmation program can reproduce S30 using the reported melt-spinning route, measure composition by EDS or WDS, verify phase constitution and determine Tc and ΔS on a common field–temperature grid [8]. A parallel S36 experiment supplies an independent fallback [5]. Identifying elemental effects requires matched Cr-free and Mn-without-Al or Al-without-Mn series under controlled processing, not simply another bootstrap of existing records. The 250–350 K application window remains primary; translated-window results guide candidate identity without assigning chemical meaning to invariant Cr or Mn–Al allocations.')
    replace(p[96],p[96].text+' Our zero-access rule and utility clipping are modeling choices. The counterexample assumes bounded candidate utilities and monotone access envelopes, not positive, convex or physically realizable complete games. Main-effect rank does not diagnose all interaction aliasing. The diagnostic AUC values and their thresholds are not externally validated across materials families. Preference regret uses an internal geometric oracle rather than independent experiments, and simple controls match Nash in this portfolio. No practical-equivalence margin or confirmatory test of benefit from bootstrap confidence was prespecified. These limitations restrict the demonstrated scope to the simulated designs and audited magnetocaloric HEA portfolio.')
    replace(p[98], 'Incomplete-support Shapley ambiguity is established in cooperative-game theory [36,37]. Our contribution is its quantitative evaluation under selective, clustered literature sampling with complete-design truth, matched random losses and conditional bootstrap confidence. A constructive pair of complete designs retains the same measured evidence but changes the Shapley leader. Veto-only comparisons do not establish a difference from random selection in either generator; they do not demonstrate equivalence. With veto and player equivalence together, direction is generator-dependent. Support diagnostics flag risk but do not certify individual chemical rankings.')
    replace(p[99], 'In the audited HEA case, 17 strategies occupy four optional-element supports. Cr veto status and Mn–Al equivalence describe access to this portfolio, not separable chemistry. S30 and S32 form the measured co-optimal set, with S36 as the independent fallback after G8 removal or restriction to H ≥ 1 T. Candidate regret remains low under the declared internal preference model, but Nash does not outperform simple deterministic controls. The reproducible workflow audits target semantics, distinguishes linear support diagnostics from Shapley identification, evaluates publication and group dependence, and keeps attribution claims separate from decisions among measured candidates. Experimental validation and broader-family transfer remain future tests.')
    replace(p[101], 'The source tables are available in the Curie-temperature auditing repository [25] and the field-normalized magnetocaloric screening repository [24]. The Curie audit record is archived at https://doi.org/10.5281/zenodo.23102726. The accompanying v6.2 reproducibility package contains audited case tables, frozen primary worlds, the expanded 600-world HEA benchmark, corrected independent-generator outputs and targeted matches, plus the verified completion witness and decision controls. A new DOI for this v6.2 package has not been assigned; the existing DOI refers to the source Curie audit, not this release. Repository links: ')
    hyperlink(p[101],'Curie-temperature audit','https://github.com/kaansarlar/Full-text-target-auditing-reveals-chemistry-regime-confounding-in-composition-only-Curie-temperature')
    p[101].add_run('; ')
    hyperlink(p[101],'magnetocaloric screening','https://github.com/kaansarlar/Machine-learning-screening-of-magnetocaloric-high-entropy-alloys-using-a-field-normalized-Sref-targ')
    replace(p[103], 'The accompanying v6.2 archive uses project-relative paths and fixed seeds. It includes scenario-level results, the two generator definitions kept separate, exact random baselines, Wilson intervals, conditional confidence denominators, editable-equation manuscript sources, integrity hashes and tests. The newly added witness and decision controls can be regenerated with scripts/run_v6_2_additions.py. The existing long benchmark outputs are frozen rather than claimed as newly simulated.')
    # Extract grid/model content for insertion using images from package, not stale rIds in another DOCX.
    grid_text=[p[78].text,p[79].text]
    predictive_text=p[48].text
    for node in [t[7]._tbl,p[49]._p,p[50]._p,p[51]._p]:node.getparent().remove(node)
    replace(p[48], 'Publication-grouped predictive gates do not support extrapolation within the intended chemistry family. The strict 3d-transition-metal Curie model gives R²=−0.336 and ρ=−0.307; the direct window classifier has an interval crossing zero and top-30% enrichment below one. Phase prediction and global entropy signals do not remove this family-transfer failure. Detailed model metrics, the full gates table and the diagnostic plot are in Supplement S2 and S10.')
    for node in [p[i]._p for i in range(77,82)]:node.getparent().remove(node)
    # Add the completion witness as the first results subsection.
    node=paragraph_after(d,p[44]._p,'3.1 Identical observed evidence permits different complete-design rankings','Heading 2')
    node=paragraph_after(d,node,'Write q(P) for the best candidate utility on support P and v(C) for the maximum q(P) over supports contained in C, including q(empty)=0. The four measured maxima are 0.164873 for Cr–Ni, 0.517081 for Cr–Ni–Cu, 0.897317 for Mn–Cr–Al and 0.947268 for Mn–Cr–Ni–Al. Completion A assigns zero to unmeasured supports. Completion B keeps all measured scores unchanged and differs only by assigning q({Mn})=0.85 to the unobserved Fe–Co-background Mn-only optional support. Both completions are bounded in [0,1] and their access envelopes are monotone. They also agree on the four observed coalition values, because 0.85 is below both measured Mn-containing maxima.')
    sh=pd.read_csv(ROOT/'results/v6_2_additions/counterexample_shapley.csv')
    table=table_after(d,node,['Player','Completion A Shapley','Completion B Shapley'],[[r.player,f'{r.phi_A:.6f}',f'{r.phi_B:.6f}'] for r in sh.itertuples()],[2.2,2.1,2.1])
    node=paragraph_after(d,table._tbl,'Table 1. Two complete bounded performance designs consistent with the same four observed coalition values and every measured strategy score. The highest allocation changes from Cr in A to Mn in B; Mn and Al cease to be equivalent in B.')
    paragraph_after(d,node,'The leader changes although the measured evidence is identical. No estimator that uses only these observations can uniquely recover both complete-design rankings. Completion A is one admissible zero-fill witness, not a justified physical imputation; B is not a proposed real alloy. This construction is a small explicit non-identification example, not a new impossibility theorem. It links the literature extension problem [36,37] to the synthetic sampled-versus-complete comparison. All 32 utilities, access values, efficiency checks and observed-value agreement tests are supplied.')
    # Put the synthetic results immediately after the witness; preserve body elements including tables.
    body=d._element.body
    nodes=list(body);start=nodes.index(p[62]._p);stop=nodes.index(p[70]._p)
    synthetic=nodes[start:stop]
    anchor=p[45]._p
    for e in synthetic:anchor.addprevious(e)
    for i,newtext in {62:'3.2 Synthetic portfolios separate stability from identification',45:'3.3 The audit changes admissible evidence',47:'3.4 Predictive models bound extrapolation',52:'3.5 Measured strategies and equal-budget controls',57:'3.6 Observed support limits element interpretation',70:'3.7 Research-group and field restrictions change the priority',75:'3.8 Temperature windows change candidate identity'}.items():replace(p[i],newtext)
    replace(p[46],'Table 3. Audited cohorts and analytical roles.')
    replace(p[58],'Table 4. The 17 eligible strategies occupy only four optional-element support classes.')
    replace(p[67],p[67].text.replace('Table 4.','Table 2.'))
    refs=['[36] Černý M and Grabisch M 2024 Incomplete cooperative games with player-centered information Discrete Appl. Math. 346 62–79 https://doi.org/10.1016/j.dam.2023.12.007',
          '[37] Černý M, Bok J, Hartman D and Hladík M 2024 Positivity and convexity in incomplete cooperative games Ann. Oper. Res. 340 785–809 https://doi.org/10.1007/s10479-024-06082-6']
    for text in refs:d.add_paragraph(text)
    format_doc(d);d.save(OUT/'MLST_HEA_manuscript_v6_2.docx')
    return predictive_text,grid_text,t[7]

def supplement(predictive_text,grid_text,gate_table):
    d=Document(ROOT/'templates/supplement_v6_1.docx');p=list(d.paragraphs)
    replace(p[0],'Supplementary information for '+TITLE);p[0].style='Title'
    replace(p[45],'Table S18. Exact-geometry comparisons with 95% Wilson recovery intervals. Each includes 0.20; this does not establish equivalence to uniform recovery. Eighteen tied-truth independent worlds are excluded from exact recovery but retained in material-error calculations.')
    replace(p[41],'Table S16. All 600 HEA-like stress-test worlds, not all of which reproduce the audited geometry. The exactly matched 358-world primary subset is reported separately in Table S18.')
    replace(p[42],'S9 Additional benchmark analyses and reproducibility map')
    replace(p[48],'The v6.2 package keeps the primary and corrected independent generator outputs separate. The primary benchmark has 4000 worlds and 500 bootstrap resamples per unit; the independent benchmark has 4000 worlds and 100 resamples, and its targeted run has 9000 worlds with 288 exact matches. Generator-specific truth, support sampling and confidence thresholds preclude pooled rates. The new completion witness and equal-budget controls are regenerated separately. Frozen source outputs are retained for provenance; they were not re-simulated for this release.')
    # Harmonize every reported primary interval with one frozen paired-bootstrap result.
    fill(d.tables[16],[row[:4]+[row[4].split(' [')[0],'['+row[4].split(' [')[1]] for row in fair_rows()])
    node=paragraph_after(d,p[30]._p,grid_text[0]+' '+grid_text[1])
    pic=d.add_paragraph();node.addnext(pic._p);pic.add_run().add_picture(str(ROOT/'figures/Figure_6_grid_stress_test.png'),width=Inches(6.4))
    paragraph_after(d,pic._p,'Figure S3. Exploratory continuous-grid game. The red star in (a) marks the highest-scoring grid composition. Numbers in (c) are omitted dataset sources S18, S36 and S42; the final bar omits S36 and S38 together. Conditional GP variance is not a calibrated total uncertainty interval.')
    replace(p[30], 'The continuous grid uses 71134 compositions, fixed fitted utilities and Gaussian-process models as described in the archived grid scripts. There are only 21 Curie and 39 entropy states. This exploratory model does not pass the empirical family-transfer gates and is not part of the complete-design identification evidence. The methods and figure are retained for transparency, not as a validated composition proposal.')
    for row in d.tables[12].rows:
        if 'Publication-bootstrap confident wrong' in row.cells[0].text:
            row.cells[1].text='45.775% of all 4000 worlds'
        if 'Group-bootstrap confident wrong' in row.cells[0].text:
            row.cells[1].text='45.575% of all 4000 worlds'
    labels={'Tc_group_block':'Tc permuted by research-group block','Tc_record':'Tc permuted by record',
            'dSref_record':'Entropy permuted by record','phase_screened':'Phase screened','primary':'Primary portfolio',
            'group bootstrap mean ci width':'Group-bootstrap mean interval width','n records':'Record count'}
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text in labels:cell.text=labels[cell.text]
    d.add_heading('S10 Detailed model diagnostics',level=1);d.add_paragraph(predictive_text)
    d._element.body.insert(len(d._element.body)-1,deepcopy(gate_table._tbl))
    d.add_paragraph('Table S21. Publication-grouped predictability gates using formula-derived primary features. These gates bound extrapolation and are not the central synthetic identification result.')
    pic=d.add_paragraph();pic.add_run().add_picture(str(ROOT/'figures/Figure_2_predictability.png'),width=Inches(6.4))
    d.add_paragraph('Figure S4. Publication-grouped validation of the audited cohorts. Global signals do not justify within-family extrapolation.')
    d.add_heading('S11 Equal-budget controls and utility-induced co-optima',level=1)
    d.add_paragraph('For 20000 Dirichlet(1,1,1) preference draws, each preference evaluates a weighted geometric mean of uT, uS and uR. The oracle maximizes this utility over the same measured candidate set. A selected list is evaluated by its best member for that preference. Ties in deterministic rule scores are resolved by original source-row order. The raw-entropy control prioritizes Tc in 250–350 K, then descending measured ΔSref; outside-window candidates are ordered by boundary distance, entropy and source order. Random controls select a size-b subset uniformly without replacement. Their expectation uses the exact probability C(j,b−1)/C(n,b) that the jth ascending utility is the subset maximum, for zero-based j. We report expected mean regret, not a quantile of averaged random regret. This internal oracle is not an independently validated model of scientific or economic preferences.')
    data=pd.read_csv(ROOT/'results/v6_2_additions/decision_controls.csv')
    rows=[]
    for rule,group in data.groupby('rule',sort=False):
        for scenario in ['All 17 eligible','Phase screened']:
            f=group[group.scenario==scenario].sort_values('budget')
            rows.append([scenario,rule]+[f'{x:.6f}' for x in f.mean_regret])
    table=table_after(d,d.paragraphs[-1]._p,['Candidate set','Decision rule','Budget 1','Budget 3','Budget 5'],rows,[1.1,2.2,1,1,1])
    paragraph_after(d,table._tbl,'Table S22. Mean regret at identical candidate budgets. Nash ties simple deterministic controls at budget one and does not establish a unique advantage. The random subset values are exact expectations conditional on each sampled preference.')
    d.add_paragraph('S30 and S32 have different raw entropy values (2.386234 and 2.074768 J kg⁻¹ K⁻¹), but both lie above the upper entropy clipping threshold. Each therefore receives temperature and clipped entropy utilities of one and reliability utility 0.85, so the equal-weight Nash score is the cube root of 0.85, or 0.947268. Their shared optimum is transformation-dependent. The single-phase screen then prioritizes S30; it does not create independent cross-group corroboration. The 108-setting sensitivity and utility-alignment null quantify this contingency rather than prove a uniquely favorable compromise.')
    d.add_heading('S12 Completion witness and inference limits',level=1)
    d.add_paragraph('The complete witness consists of results/v6_2_additions/counterexample_all_32_supports.csv and counterexample_shapley.csv. The script verifies equality at all measured supports, equality of access values for those four supports, bounded utilities, monotone access envelopes and Shapley efficiency. One additional optional Mn-only support in completion B breaks the observed Mn–Al symmetry and changes the leading player. The construction is not constrained to additive, positive or convex games and is not claimed physically realizable. It establishes non-identification under the bounded-access-game class, not a new theorem or chemical mechanism. See main references [36,37] for incomplete-game extension theory.')
    d.add_paragraph('All paired comparison intervals are taken from the same frozen fairness summary. Failure to exclude zero is not an equivalence test; failure to exclude 0.20 does not show uniform accuracy. No practical-equivalence margin was prespecified. Main-effect rank is a diagnostic for a linear model and is insufficient for interaction or complete-game identification. Differences in conditional confidence subsets do not constitute a causal or paired test of benefit from confidence.')
    format_doc(d);d.save(OUT/'MLST_HEA_supplement_v6_2.docx')

def checks():
    import zipfile
    d=Document(OUT/'MLST_HEA_manuscript_v6_2.docx')
    text='\n'.join(p.text for p in d.paragraphs)
    for bad in ['random-level','indistinguishable','identified in version 5','This is a Nash bargaining solution']:
        assert bad not in text,bad
    with zipfile.ZipFile(OUT/'MLST_HEA_manuscript_v6_2.docx') as z:
        xml=z.read('word/document.xml').decode()
        equations=len(re.findall(r'<m:oMath[ >]',xml))
    with zipfile.ZipFile(ROOT/'templates/manuscript_v6_1.docx') as z:
        original_equations=len(re.findall(r'<m:oMath[ >]',z.read('word/document.xml').decode()))
    assert equations==original_equations==6,equations
    citations=set()
    for p in d.paragraphs:
        if re.match(r'^\[\d+\]',p.text):continue
        for group in re.findall(r'\[([\d,\-– ]+)\]',p.text):
            group=group.replace('–','-')
            for item in group.split(','):
                if '-' in item:
                    a,b=item.split('-');citations.update(range(int(a),int(b)+1))
                elif item.strip():citations.add(int(item))
    assert set(range(1,38))<=citations,sorted(set(range(1,38))-citations)
    report={'native_Word_equations':equations,'main_figures':len(d.inline_shapes),'all_37_references_cited':True,'text_checks_passed':True}
    (ROOT/'results/v6_2_additions/document_checks.json').write_text(json.dumps(report,indent=2))
    print(report)

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    values=main_document();supplement(*values);checks()
